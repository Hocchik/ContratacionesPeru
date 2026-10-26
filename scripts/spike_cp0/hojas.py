"""Lista hojas, dimensiones y encabezados de cada .xlsx sin modificarlo."""
import sys
from pathlib import Path

from openpyxl import load_workbook

for ruta in sys.argv[1:]:
    wb = load_workbook(Path(ruta), read_only=True)
    print(f"== {Path(ruta).name}")
    for ws in wb.worksheets:
        filas = list(ws.iter_rows(min_row=1, max_row=3, values_only=True))
        print(f"  hoja={ws.title!r} max_row={ws.max_row} max_col={ws.max_column}")
        for i, f in enumerate(filas, 1):
            print(f"   fila{i}: {[c for c in f][:40]}")
    wb.close()
