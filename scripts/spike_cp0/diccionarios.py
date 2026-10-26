"""Vuelca el contenido de los diccionarios (.xlsx y .docx) como texto, sin modificarlos."""
import re
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook

sys.stdout.reconfigure(encoding="utf-8")

for ruta in sorted(Path(sys.argv[1]).iterdir()):
    print(f"\n######## {ruta.name}")
    if ruta.suffix.lower() == ".xlsx":
        wb = load_workbook(ruta, read_only=True, data_only=True)
        for ws in wb.worksheets:
            print(f"--- hoja {ws.title!r}")
            for fila in ws.iter_rows(values_only=True):
                celdas = [str(c).strip().replace("\n", " ") for c in fila if c is not None and str(c).strip()]
                if celdas:
                    print(" | ".join(celdas))
        wb.close()
    elif ruta.suffix.lower() == ".docx":
        with zipfile.ZipFile(ruta) as z:
            xml = z.read("word/document.xml").decode("utf-8")
        parrafos = [re.sub(r"<[^>]+>", "", p) for p in re.split(r"</w:p>|</w:tc>", xml)]
        print("\n".join(p.strip() for p in parrafos if p.strip()))
